"""Lê planilhas mensais da Ouvidoria (layout das 35 colunas) e gera o SQL que grava as linhas
na área de carga public.importacao_linhas. Depois, no banco:

    select * from public.importacao_processar('<lote>');        -- confere, sem gravar nada
    select aba, linha, erros, avisos from public.importacao_linhas where lote = '<lote>' and erros <> '{}';
    select * from public.importacao_processar('<lote>', true);  -- importa as linhas sem erro

Uso: python3 scripts/importar_planilha.py <lote> planilha1.xlsx [planilha2.xlsx ...] > carga.sql
"""
import datetime
import json
import re
import sys

import openpyxl

# Colunas A a AI da planilha, na ordem.
CAMPOS = [
    'n', 'situacao', 'protocolo', 'entrada', 'falabr', 'tipo', 'identificacao', 'cadastro', 'recebimento',
    'assunto', 'logradouro', 'numero', 'complemento', 'bairro', 'cep', 'descricao', 'diretoria', 'avaliador',
    'encaminhado', 'prazo', 'prazo_prorrogacao', 'resposta_data', 'tempo', 'resposta', 'resolutiva',
    'arquivamento', 'reab_pedido', 'reab_texto', 'reab_aceita', 'reab_data', 'reab_prazo',
    'reab_prazo_prorrogacao', 'reab_resposta_data', 'reab_resposta', 'observacoes',
]
DATAS = {'cadastro', 'recebimento', 'encaminhado', 'prazo', 'prazo_prorrogacao', 'resposta_data',
         'reab_pedido', 'reab_data', 'reab_prazo', 'reab_prazo_prorrogacao', 'reab_resposta_data'}
# Campos que, todos vazios, indicam linha em branco.
ESSENCIAIS = ['protocolo', 'falabr', 'entrada', 'tipo', 'assunto', 'descricao', 'cadastro']


def valor(campo, celula):
    if celula is None:
        return None
    if isinstance(celula, datetime.datetime):
        return celula.strftime('%Y-%m-%d %H:%M') if (celula.hour or celula.minute) else celula.strftime('%Y-%m-%d')
    if isinstance(celula, datetime.date):
        return celula.isoformat()
    if isinstance(celula, float) and celula.is_integer():
        celula = int(celula)
    texto = str(celula).strip()
    if not texto:
        return None
    if campo in DATAS:
        # "00/00/0000" é o marcador de data vazia do modelo da planilha.
        if re.fullmatch(r'[0/\s:]+', texto):
            return None
        achado = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})(?:\s+(\d{1,2}):(\d{2}))?', texto)
        if achado:
            dia, mes, ano, hora, minuto = achado.groups()
            try:
                data = datetime.datetime(int(ano), int(mes), int(dia), int(hora or 0), int(minuto or 0))
                return valor(campo, data)
            except ValueError:
                return texto
    return texto


def linhas(caminho):
    livro = openpyxl.load_workbook(caminho, data_only=True)
    for aba in livro.worksheets:
        cabecalho = None
        for i, linha in enumerate(aba.iter_rows(min_row=1, max_row=20, max_col=1, values_only=True), 1):
            if str(linha[0] or '').strip().upper() in ('Nº', 'N°', 'NO', 'N'):
                cabecalho = i
                break
        if cabecalho is None:
            continue  # aba de listas (CAMPOS) ou fora do layout
        # A linha logo abaixo do cabeçalho é a de ajuda; os dados começam depois dela.
        for i, celulas in enumerate(aba.iter_rows(min_row=cabecalho + 2, max_col=len(CAMPOS), values_only=True),
                                    cabecalho + 2):
            dados = {campo: valor(campo, c) for campo, c in zip(CAMPOS, celulas)}
            if all(dados.get(c) is None for c in ESSENCIAIS):
                continue
            yield f'{caminho.rsplit("/", 1)[-1]} › {aba.title}', i, {k: v for k, v in dados.items() if v is not None}


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    lote = sys.argv[1]
    n = 0
    print('begin;')
    for caminho in sys.argv[2:]:
        for aba, linha, dados in linhas(caminho):
            corpo = json.dumps(dados, ensure_ascii=False)
            print('insert into public.importacao_linhas (lote, aba, linha, dados) values '
                  f"({sql(lote)}, {sql(aba)}, {linha}, {sql(corpo)}::jsonb) on conflict (lote, aba, linha) do nothing;")
            n += 1
    print('commit;')
    print(f'-- {n} linhas', file=sys.stderr)


def sql(texto):
    return "'" + texto.replace("'", "''") + "'"


if __name__ == '__main__':
    main()
